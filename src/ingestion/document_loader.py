"""
Document loaders for various file formats.
"""
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
import openpyxl
from pathlib import Path
from typing import Dict, List, Any
import logging

logger = logging.getLogger(__name__)


class DocumentLoader:
    """Factory for loading different document types."""
    
    @staticmethod
    def load(file_path: str) -> Dict[str, Any]:
        """Load document based on file extension."""
        path = Path(file_path)
        extension = path.suffix.lower()
        
        loaders = {
            '.pdf': PDFLoader,
            '.docx': DOCXLoader,
            '.pptx': PPTXLoader,
            '.xlsx': ExcelLoader,
            '.txt': TextLoader
        }
        
        loader_class = loaders.get(extension)
        if not loader_class:
            raise ValueError(f"Unsupported file type: {extension}")
        
        return loader_class.load(file_path)


class PDFLoader:
    """Load PDF documents."""
    
    @staticmethod
    def load(file_path: str) -> Dict[str, Any]:
        """Extract text and metadata from PDF."""
        try:
            doc = fitz.open(file_path)
            chunks = []
            
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                
                if text.strip():
                    chunks.append({
                        'text': text,
                        'metadata': {
                            'source': Path(file_path).name,
                            'page': page_num + 1,
                            'total_pages': len(doc),
                            'file_type': 'pdf'
                        }
                    })
            
            doc.close()
            
            return {
                'chunks': chunks,
                'metadata': {
                    'source': Path(file_path).name,
                    'total_pages': len(doc),
                    'file_type': 'pdf'
                }
            }
        
        except Exception as e:
            logger.error(f"Error loading PDF {file_path}: {e}")
            raise


class DOCXLoader:
    """Load DOCX documents."""
    
    @staticmethod
    def load(file_path: str) -> Dict[str, Any]:
        """Extract text and metadata from DOCX."""
        try:
            doc = Document(file_path)
            full_text = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            
            text = '\n\n'.join(full_text)
            
            return {
                'chunks': [{
                    'text': text,
                    'metadata': {
                        'source': Path(file_path).name,
                        'file_type': 'docx',
                        'paragraphs': len(full_text)
                    }
                }],
                'metadata': {
                    'source': Path(file_path).name,
                    'file_type': 'docx',
                    'paragraphs': len(full_text)
                }
            }
        
        except Exception as e:
            logger.error(f"Error loading DOCX {file_path}: {e}")
            raise


class PPTXLoader:
    """Load PowerPoint documents."""
    
    @staticmethod
    def load(file_path: str) -> Dict[str, Any]:
        """Extract text and metadata from PPTX."""
        try:
            prs = Presentation(file_path)
            chunks = []
            
            for slide_num, slide in enumerate(prs.slides):
                slide_text = []
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)
                
                if slide_text:
                    chunks.append({
                        'text': '\n'.join(slide_text),
                        'metadata': {
                            'source': Path(file_path).name,
                            'slide': slide_num + 1,
                            'total_slides': len(prs.slides),
                            'file_type': 'pptx'
                        }
                    })
            
            return {
                'chunks': chunks,
                'metadata': {
                    'source': Path(file_path).name,
                    'total_slides': len(prs.slides),
                    'file_type': 'pptx'
                }
            }
        
        except Exception as e:
            logger.error(f"Error loading PPTX {file_path}: {e}")
            raise


class ExcelLoader:
    """Load Excel documents."""
    
    @staticmethod
    def load(file_path: str) -> Dict[str, Any]:
        """Extract text and metadata from Excel."""
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            chunks = []
            
            for sheet in wb.worksheets:
                rows_text = []
                
                for row in sheet.iter_rows(values_only=True):
                    row_text = [str(cell) for cell in row if cell is not None]
                    if row_text:
                        rows_text.append(' | '.join(row_text))
                
                if rows_text:
                    chunks.append({
                        'text': '\n'.join(rows_text),
                        'metadata': {
                            'source': Path(file_path).name,
                            'sheet': sheet.title,
                            'file_type': 'xlsx'
                        }
                    })
            
            return {
                'chunks': chunks,
                'metadata': {
                    'source': Path(file_path).name,
                    'sheets': [sheet.title for sheet in wb.worksheets],
                    'file_type': 'xlsx'
                }
            }
        
        except Exception as e:
            logger.error(f"Error loading Excel {file_path}: {e}")
            raise


class TextLoader:
    """Load plain text documents."""
    
    @staticmethod
    def load(file_path: str) -> Dict[str, Any]:
        """Extract text from plain text file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            
            return {
                'chunks': [{
                    'text': text,
                    'metadata': {
                        'source': Path(file_path).name,
                        'file_type': 'txt'
                    }
                }],
                'metadata': {
                    'source': Path(file_path).name,
                    'file_type': 'txt'
                }
            }
        
        except Exception as e:
            logger.error(f"Error loading text file {file_path}: {e}")
            raise
